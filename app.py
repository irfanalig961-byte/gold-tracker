import io
from datetime import date

import pandas as pd
import plotly.graph_objects as go
import requests
import streamlit as st
import yfinance as yf
from plotly.subplots import make_subplots

# ---------------- settings: the only lines you may ever need to change ----------------
START_DATE = "2026-09-29"
NOTES_URL = "https://docs.google.com/spreadsheets/d/e/2PACX-1vSPU8Q9EaMP1vNOMkcKnnwo5zs3XXyubtLbVWy3__SMqE1FZ1pnZhegHEa6xTzIE2e2MQai1CWmpqjk/pub?gid=0&single=true&output=csv"
# ----------------------------------------------------------------------------------------

COLS = ["date", "title", "step1", "step2", "step3", "gold_source", "gold_link", "dollar_note", "dollar_source", "dollar_link", "silver_note", "silver_source", "silver_link"]

st.set_page_config(page_title="Gold Tracker", page_icon="logo.png", layout="wide")
st.logo("logo.png", size="large")


@st.cache_data(ttl=3600)
def load_prices(start):
    tickers = {"GC=F": "Gold", "SI=F": "Silver", "DX-Y.NYB": "Dollar"}
    data = yf.download(list(tickers), start=start, progress=False)["Close"]
    data = data.rename(columns=tickers)[["Gold", "Dollar", "Silver"]].dropna()
    data.index = pd.to_datetime(data.index).tz_localize(None).normalize()
    return data


@st.cache_data(ttl=300)
def load_notes(url):
    r = requests.get(url, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
    if r.status_code != 200: raise ValueError(f"Google answered with error {r.status_code}.")
    text = r.content.decode("utf-8-sig")
    if text.lstrip().startswith("<"): raise ValueError("Google sent a web page instead of the notes.")
    notes = pd.read_csv(io.StringIO(text), dtype=str).fillna("")
    notes.columns = [c.strip().lower() for c in notes.columns]
    for col in [c for c in COLS if c not in notes.columns]: notes[col] = ""
    notes["date"] = pd.to_datetime(notes["date"].str.strip(), errors="coerce").dt.normalize()
    return notes.dropna(subset=["date"]).sort_values("date")


def pct(x):
    return "" if pd.isna(x) else f"{x:+.2f}%"


def news_line(icon, label, note, source, link):
    text = f"{icon} {label}: {note}" if note else f"{icon} {label} news:"
    return text + (f" [{source or 'read the article'}]({link})" if link else "")


# ---------------- data ----------------
try:
    data = load_prices(START_DATE)
except Exception as e:
    st.error(f"Could not download prices from Yahoo Finance just now. Refresh in a minute. ({e})")
    st.stop()
if data.empty: st.error("No prices came back from Yahoo Finance. Refresh in a minute."); st.stop()

notes_error = None
try:
    notes = load_notes(NOTES_URL)
except Exception as e:
    notes_error = str(e)
    notes = pd.DataFrame(columns=COLS)
    notes["date"] = pd.to_datetime(notes["date"])

moves = data.pct_change() * 100
notes_in_data = notes[notes["date"].isin(data.index)]

# ---------------- header ----------------
st.title("Gold Tracker")
st.caption("Gold next to the asset that usually moves against it (the US dollar) " "and the one that usually moves with it (silver). Click any day on the chart to see what happened.")

last = data.index[-1]
c1, c2, c3 = st.columns(3)
c1.metric("Gold (US$/oz)", f"${data['Gold'].iloc[-1]:,.0f}", pct(moves["Gold"].iloc[-1]))
c2.metric("US dollar index", f"{data['Dollar'].iloc[-1]:.2f}", pct(moves["Dollar"].iloc[-1]))
c3.metric("Silver (US$/oz)", f"${data['Silver'].iloc[-1]:.2f}", pct(moves["Silver"].iloc[-1]))
st.caption(f"Latest close: {last:%a %d %b %Y}. Prices refresh hourly, notes every 5 minutes. " f"Notes loaded: {len(notes)}.")
if notes_error: st.warning(f"Could not read the notes sheet: {notes_error}")
if not notes_error and len(notes) == 0: st.info("The notes sheet was read, but no rows had a valid date (e.g. 2026-10-09).")
if st.button("Reload notes now"): load_notes.clear(); st.rerun()

with st.expander("What is measured, and how to read it"):
    st.markdown("""
- **Gold**: COMEX gold futures (GC=F), US dollars per troy ounce, daily closing price. Gold line.
- **Silver**: COMEX silver futures (SI=F), US dollars per troy ounce, daily closing price. Blue line.
- **US dollar index** (DX-Y.NYB): the dollar against six currencies, weighted euro 57.6%, Japanese yen 13.6%,
  British pound 11.9%, Canadian dollar 9.1%, Swedish krona 4.2%, Swiss franc 3.6%. Green line.
- **Black dots**: days with a written note for that line.
- **Daily % change**: today's close compared with yesterday's close.
- **Correlation**: how closely daily % changes move with gold, using all days so far.
  −1 = always opposite, 0 = no link, +1 = always together.
  The dollar is usually negative with gold, silver usually positive.
- Source: Yahoo Finance (prices) and my own notes (news).
""")

# ---------------- chart ----------------
corr_dollar = moves["Gold"].expanding(min_periods=4).corr(moves["Dollar"])
corr_silver = moves["Gold"].expanding(min_periods=4).corr(moves["Silver"])

fig = make_subplots(rows=4, cols=1, shared_xaxes=True, vertical_spacing=0.07, row_heights=[0.34, 0.2, 0.2, 0.26], subplot_titles=["Gold, US$ per troy ounce (black dots = days with notes)", "US dollar index: usually moves OPPOSITE to gold", "Silver, US$ per troy ounce: usually moves WITH gold", "Correlation with gold so far (−1 opposite · 0 no link · +1 together)"])
fig.add_trace(go.Scatter(x=data.index, y=data["Gold"], name="Gold", mode="lines+markers", marker=dict(size=6), line=dict(color="goldenrod")), row=1, col=1)
fig.add_trace(go.Scatter(x=data.index, y=data["Dollar"], name="Dollar", line=dict(color="seagreen")), row=2, col=1)
fig.add_trace(go.Scatter(x=data.index, y=data["Silver"], name="Silver", line=dict(color="steelblue")), row=3, col=1)
fig.add_trace(go.Scatter(x=corr_dollar.index, y=corr_dollar, name="Dollar vs gold", line=dict(color="seagreen")), row=4, col=1)
fig.add_trace(go.Scatter(x=corr_silver.index, y=corr_silver, name="Silver vs gold", line=dict(color="steelblue")), row=4, col=1)
fig.add_hline(y=0, line_color="grey", row=4, col=1)
fig.update_yaxes(range=[-1, 1], title_text="−1 to +1", row=4, col=1)
fig.update_yaxes(title_text="US$/oz", row=1, col=1)
fig.update_yaxes(title_text="Index", row=2, col=1)
fig.update_yaxes(title_text="US$/oz", row=3, col=1)

# black dots: days with a note, on each line that has news for that day
gold_days = notes_in_data[(notes_in_data["title"] != "") | (notes_in_data["gold_link"] != "")]
dollar_days = notes_in_data[(notes_in_data["dollar_note"] != "") | (notes_in_data["dollar_link"] != "")]
silver_days = notes_in_data[(notes_in_data["silver_note"] != "") | (notes_in_data["silver_link"] != "")]
fig.add_trace(go.Scatter(x=gold_days["date"], y=data["Gold"].reindex(gold_days["date"]), mode="markers", marker=dict(size=11, color="black"), text=gold_days["title"], hoverinfo="text", name="Gold news", showlegend=False), row=1, col=1)
fig.add_trace(go.Scatter(x=dollar_days["date"], y=data["Dollar"].reindex(dollar_days["date"]), mode="markers", marker=dict(size=11, color="black"), text=dollar_days["dollar_note"], hoverinfo="text", name="Dollar news", showlegend=False), row=2, col=1)
fig.add_trace(go.Scatter(x=silver_days["date"], y=data["Silver"].reindex(silver_days["date"]), mode="markers", marker=dict(size=11, color="black"), text=silver_days["silver_note"], hoverinfo="text", name="Silver news", showlegend=False), row=3, col=1)
fig.update_layout(height=900, hovermode="x unified", clickmode="event+select", legend=dict(orientation="h", y=-0.05), margin=dict(t=40, b=10))

event = st.plotly_chart(fig, width="stretch", on_select="rerun", selection_mode="points", key="chart")

# which day did the viewer pick?
try:
    pts = event.selection.points
    picked = pd.Timestamp(str(pts[0]["x"])[:10]) if pts else None
except Exception:
    picked = None

all_days = list(data.index[::-1])
note_days = set(notes_in_data["date"])
default = notes_in_data["date"].iloc[-1] if len(notes_in_data) else data.index[-1]
default = picked if picked is not None and picked in data.index else default
day = st.selectbox("Day", all_days, index=all_days.index(default), format_func=lambda d: d.strftime("%a %d %b %Y") + ("  ●" if d in note_days else ""))

# ---------------- day file ----------------
left, right = st.columns([3, 2])
row = notes_in_data[notes_in_data["date"] == day]
n = row.iloc[-1] if len(row) else None
left.subheader(day.strftime("%A %d %B %Y"))
if n is None: left.info("No note for this day yet.")
if n is not None: left.markdown(f"#### {n['title']}")
if n is not None: left.markdown("\n".join(f"{i}. {t}" for i, t in enumerate([t for t in [n["step1"], n["step2"], n["step3"]] if t.strip()], 1)))
if n is not None and n["gold_link"]: left.markdown(news_line("📰", "Gold", "", n["gold_source"], n["gold_link"]))
if n is not None and (n["dollar_note"] or n["dollar_link"]): left.markdown(news_line("💵", "Dollar", n["dollar_note"], n["dollar_source"], n["dollar_link"]))
if n is not None and (n["silver_note"] or n["silver_link"]): left.markdown(news_line("🥈", "Silver", n["silver_note"], n["silver_source"], n["silver_link"]))

right.markdown("**Prices that day**")
right.dataframe(pd.DataFrame({"Close": data.loc[day].round(2), "Day change": moves.loc[day].map(pct)}), width="stretch")
g, x, s = moves.loc[day, "Gold"], moves.loc[day, "Dollar"], moves.loc[day, "Silver"]
check = bool(pd.notna(g) and abs(g) >= 0.05)
if check: right.write("✅ Dollar moved opposite to gold." if g * x < 0 else "❌ Dollar moved the same way as gold.")
if check: right.write("✅ Silver moved with gold." if g * s > 0 else "❌ Silver moved against gold.")

# ---------------- correlation summary ----------------
st.divider()
m = moves.dropna()
a, b = st.columns(2)
a.metric("Dollar vs gold, correlation so far", f"{m['Gold'].corr(m['Dollar']):+.2f}", help="−1 = always opposite, +1 = always together")
b.metric("Silver vs gold, correlation so far", f"{m['Gold'].corr(m['Silver']):+.2f}", help="−1 = always opposite, +1 = always together")
st.caption(f"Based on {len(m)} trading days. Fewer than 20 days is only a rough guide.")

# ---------------- all notes ----------------
st.subheader("All notes")
show = notes.copy()
show["date"] = show["date"].dt.date
if len(show) == 0: st.write("No notes yet.")
if len(show) > 0: st.dataframe(show[["date", "title", "step1", "step2", "step3", "gold_link", "dollar_note", "dollar_link", "silver_note", "silver_link"]], width="stretch", hide_index=True, column_config={"gold_link": st.column_config.LinkColumn("Gold news"), "dollar_link": st.column_config.LinkColumn("Dollar news"), "silver_link": st.column_config.LinkColumn("Silver news")})


# ---------------- Excel downloads: a readable report + the raw data ----------------
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INK, GREY, GOLDC, UP, DOWN, CARD, LINE = "222222", "777777", "8B6914", "1E7B34", "C0392B", "FBF6E9", "E6DCC3"


def _font(size=11, bold=False, color=INK, italic=False, underline=None):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic, underline=underline)


def _say(ws, cell, text, **f):
    ws[cell] = text
    ws[cell].font = _font(**f)
    ws[cell].alignment = Alignment(vertical="top", wrap_text=True)


def _block(ws, row, text, height_per_line=15, chars_per_line=115, **f):
    # one wide wrapped paragraph across columns B:I
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
    _say(ws, f"B{row}", text, **f)
    lines = max(1, -(-len(text) // chars_per_line))
    ws.row_dimensions[row].height = lines * height_per_line + 4


def _signed(x, suffix="%"):
    return "–" if pd.isna(x) else f"{x:+.1f}{suffix}"


def _chart(ws, src, col, title, anchor, rows):
    ch = LineChart()
    ch.title = title
    ch.legend = None
    ch.height, ch.width = 6.2, 8.2
    ch.add_data(Reference(src, min_col=col, min_row=1, max_row=rows + 1), titles_from_data=True)
    ch.set_categories(Reference(src, min_col=1, min_row=2, max_row=rows + 1))
    vals = [v for v in (src.cell(r, col).value for r in range(2, rows + 2)) if isinstance(v, (int, float))]
    if vals:
        pad = (max(vals) - min(vals)) * 0.15 or max(vals) * 0.01
        ch.y_axis.scaling.min = round(min(vals) - pad, 2)
        ch.y_axis.scaling.max = round(max(vals) + pad, 2)
    ch.x_axis.number_format = "dd mmm"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines = None
    ch.series[0].graphicalProperties.line.solidFill = {"Gold": "B8860B", "Dollar": "2E8B57", "Silver": "4682B4"}[title.split()[0]]
    ch.series[0].graphicalProperties.line.width = 28000
    ch.series[0].smooth = False
    ws.add_chart(ch, anchor)


def _data_sheet(ws, df, index_label, pct_word):
    head = [index_label] + list(df.columns)
    ws.append(head)
    for idx, row in df.iterrows():
        ws.append([idx] + [None if pd.isna(v) else float(v) for v in row])
    for c, h in enumerate(head, 1):
        cell = ws.cell(1, c)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GOLDC)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = 13 if c > 1 else 14
        for r in range(2, ws.max_row + 1):
            x = ws.cell(r, c)
            x.font = _font()
            if c == 1:
                x.number_format = "ddd dd mmm yyyy"
            elif pct_word in h:
                x.number_format = '+0.00"%";-0.00"%";0.00"%"'
                if isinstance(x.value, float) and x.value != 0:
                    x.font = _font(color=UP if x.value > 0 else DOWN)
            else:
                x.number_format = "#,##0.00"
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"


def _events_sheet(ws, ev):
    cols = [("date", "Date", 13), ("title", "Headline", 34), ("step1", "What happened 1", 40), ("step2", "What happened 2", 40),
            ("step3", "What happened 3", 40), ("gold_link", "Gold news", 12), ("dollar_note", "Dollar", 40),
            ("dollar_link", "Dollar news", 12), ("silver_note", "Silver", 40), ("silver_link", "Silver news", 12)]
    for c, (_, label, w) in enumerate(cols, 1):
        cell = ws.cell(1, c, label)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GOLDC)
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    for r, (_, n) in enumerate(ev.iterrows(), 2):
        longest = 1
        for c, (key, _, w) in enumerate(cols, 1):
            v = n.get(key, "")
            cell = ws.cell(r, c)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = _font(bold=(key == "title"))
            if key == "date":
                cell.value = pd.Timestamp(v).date()
                cell.number_format = "ddd dd mmm"
            elif key.endswith("_link"):
                if str(v).startswith("http"):
                    src = n.get(key.replace("_link", "_source"), "") or "Open"
                    cell.value = f"{src} ↗"
                    cell.hyperlink = str(v)
                    cell.font = _font(color="1F5FAD", underline="single")
            else:
                cell.value = v
                longest = max(longest, -(-len(str(v)) // int(w * 1.1)))
        ws.row_dimensions[r].height = longest * 14 + 6
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "B2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"


def excel_bytes(px, ev, title, weekly=None):
    # px: prices for the period (Gold, Dollar, Silver), ev: notes for the period
    px = px.dropna(how="all")
    before = data[data.index < px.index[0]]
    base = before.iloc[-1] if len(before) else px.iloc[0]          # close just before the period
    mv = moves.reindex(px.index)                                    # daily % moves from the full history
    wb = Workbook()
    rep = wb.active
    rep.title = "Report"
    raw = wb.create_sheet("Daily prices")
    table = px.round(2).join(mv.round(2).add_suffix(" % change"))
    table.index = table.index.date
    _data_sheet(raw, table, "Date", "%")
    _events_sheet(wb.create_sheet("Events"), ev.sort_values("date", ascending=False))
    if weekly is not None:
        _data_sheet(wb.create_sheet("Weekly"), weekly, "Week ending", "%")

    # ----- the report page -----
    rep.sheet_view.showGridLines = False
    rep.column_dimensions["A"].width = 3
    for c in "BCDEFGHI":
        rep.column_dimensions[c].width = 13.5
    first, lastd = px.index[0], px.index[-1]
    _block(rep, 2, title, size=22, bold=True, height_per_line=30)
    _block(rep, 3, f"{first:%a %d %b %Y} to {lastd:%a %d %b %Y}  ·  {len(px)} trading days  ·  {len(ev)} notes", color=GREY)

    # three price cards
    edge = Side(style="thin", color=LINE)
    cards = [("Gold", "US$ per troy ounce", "B", "C", "{:,.0f}"), ("Dollar", "US dollar index", "E", "F", "{:.2f}"),
             ("Silver", "US$ per troy ounce", "H", "I", "{:.2f}")]
    for name, unit, c1, c2, fmt in cards:
        chg = (px[name].iloc[-1] / base[name] - 1) * 100
        for r in range(5, 9):
            rep.merge_cells(f"{c1}{r}:{c2}{r}")
            for c in (c1, c2):
                rep[f"{c}{r}"].fill = PatternFill("solid", fgColor=CARD)
                rep[f"{c}{r}"].border = Border(left=edge if c == c1 else None, right=edge if c == c2 else None,
                                                top=edge if r == 5 else None, bottom=edge if r == 8 else None)
            rep[f"{c1}{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        rep[f"{c1}5"] = name.upper()
        rep[f"{c1}5"].font = _font(size=9, bold=True, color=GREY)
        rep[f"{c1}6"] = fmt.format(px[name].iloc[-1])
        rep[f"{c1}6"].font = _font(size=24, bold=True)
        rep[f"{c1}7"] = f"{_signed(chg)} this period"
        rep[f"{c1}7"].font = _font(size=11, bold=True, color=UP if chg > 0 else DOWN if chg < 0 else GREY)
        rep[f"{c1}8"] = unit
        rep[f"{c1}8"].font = _font(size=9, color=GREY)
    rep.row_dimensions[6].height = 34

    # how they moved together, in plain words
    _block(rep, 10, "How they moved together", size=14, bold=True, height_per_line=22)
    m = mv.dropna()
    days = len(m)
    opp = int(((m["Gold"] * m["Dollar"]) < 0).sum())
    same = int(((m["Gold"] * m["Silver"]) > 0).sum())
    cd = m["Gold"].corr(m["Dollar"]) if days > 2 else float("nan")
    cs = m["Gold"].corr(m["Silver"]) if days > 2 else float("nan")
    lines = [
        f"Dollar vs gold: moved in opposite directions on {opp} of {days} days"
        + ("" if pd.isna(cd) else f" (correlation {cd:+.2f}). ") + ("The usual pattern held." if opp > days / 2 else "The usual pattern did NOT hold this period."),
        f"Silver vs gold: moved in the same direction on {same} of {days} days"
        + ("" if pd.isna(cs) else f" (correlation {cs:+.2f}). ") + ("The usual pattern held." if same > days / 2 else "The usual pattern did NOT hold this period."),
        "Correlation runs from −1 (always opposite) to +1 (always together). With fewer than 20 days it is only a rough guide.",
    ]
    for i, t in enumerate(lines):
        _block(rep, 11 + i, t, color=GREY if i == 2 else INK, italic=(i == 2), size=10 if i == 2 else 11)

    # small charts
    _block(rep, 15, "Prices", size=14, bold=True, height_per_line=22)
    n = len(px)
    _chart(rep, raw, 2, "Gold (US$/oz)", "B16", n)
    _chart(rep, raw, 3, "Dollar index", "E16", n)
    _chart(rep, raw, 4, "Silver (US$/oz)", "H16", n)

    # the story, day by day, newest first
    r = 29
    _block(rep, r, "What happened, day by day", size=14, bold=True, height_per_line=22)
    r += 1
    if len(ev) == 0:
        _block(rep, r, "No notes for this period yet.", color=GREY)
    for _, nt in ev.sort_values("date", ascending=False).iterrows():
        d = pd.Timestamp(nt["date"])
        r += 1
        rep.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        _say(rep, f"B{r}", f"{d:%A %d %B}", size=10, bold=True, color=GOLDC)
        if d in mv.index:
            mvd = mv.loc[d]
            rep.merge_cells(start_row=r, start_column=5, end_row=r, end_column=9)
            _say(rep, f"E{r}", f"Gold {_signed(mvd['Gold'])}   ·   Dollar {_signed(mvd['Dollar'])}   ·   Silver {_signed(mvd['Silver'])}", size=10, color=GREY)
            rep[f"E{r}"].alignment = Alignment(horizontal="right")
        for c in "BCDEFGHI":
            rep[f"{c}{r}"].border = Border(top=Side(style="thin", color=LINE))
        r += 1
        _block(rep, r, nt["title"], size=13, bold=True, height_per_line=18, chars_per_line=85)
        steps = [s for s in (nt["step1"], nt["step2"], nt["step3"]) if str(s).strip()]
        for i, s in enumerate(steps, 1):
            r += 1
            _block(rep, r, f"{i}.  {s}")
        for label, key in (("Dollar", "dollar_note"), ("Silver", "silver_note")):
            if str(nt.get(key, "")).strip():
                r += 1
                _block(rep, r, f"{label}: {nt[key]}", color="444444")
        links = [(f"{(nt.get(k + '_source') or k.title())} ({k}) ↗", nt.get(k + "_link", "")) for k in ("gold", "dollar", "silver")]
        links = [(t, u) for t, u in links if str(u).startswith("http")]
        if links:
            r += 1
            for (t, u), c in zip(links, ("B", "E", "H")):
                c2 = chr(ord(c) + 2)
                rep.merge_cells(f"{c}{r}:{c2}{r}")
                rep[f"{c}{r}"] = t
                rep[f"{c}{r}"].hyperlink = str(u)
                rep[f"{c}{r}"].font = _font(size=10, color="1F5FAD", underline="single")
        r += 1
        rep.row_dimensions[r].height = 10
    _block(rep, r + 1, "Sources: prices from Yahoo Finance (COMEX gold and silver futures, ICE US dollar index); news notes are my own, with links to the original articles.", size=9, color=GREY, italic=True)

    rep.page_setup.orientation = "portrait"
    rep.page_setup.fitToWidth = 1
    rep.page_setup.fitToHeight = 0
    rep.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


weekly = data.resample("W-FRI").last()
weekly = weekly.round(2).join((weekly.pct_change() * 100).round(2).add_suffix(" % week"))
weekly.index = weekly.index.date

week_start = last - pd.Timedelta(days=last.weekday())
week_px = data[data.index >= week_start]
week_ev = notes[notes["date"] >= week_start]
d1, d2 = st.columns(2)
d1.download_button("Download everything (Excel)", excel_bytes(data, notes, "Gold Tracker report", weekly), file_name=f"gold_tracker_{date.today().isoformat()}.xlsx")
d2.download_button("Download this week (Excel)", excel_bytes(week_px, week_ev, f"Gold Tracker: week of {week_start:%d %b %Y}"), file_name=f"gold_week_of_{week_start:%Y-%m-%d}.xlsx")
